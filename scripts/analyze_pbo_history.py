"""Aggregate the complete PBO workbook series from fiscal years 2558 to 2568."""

from __future__ import annotations

import argparse
import json
import pathlib
import sys
from collections import defaultdict
from itertools import chain

import openpyxl


THEMES = {
    "ict": ("คอมพิวเตอร์", "ซอฟต์แวร์", "ดิจิทัล", "สารสนเทศ", "ระบบเครือข่าย", "ฐานข้อมูล"),
    "training": ("ฝึกอบรม", "อบรม", "สัมมนา", "ศึกษาดูงาน"),
    "land": ("ก่อสร้าง", "ปรับปรุง", "ที่ดิน", "อาคาร", "ถนน", "สะพาน"),
    "procurement": ("จัดซื้อ", "จัดจ้าง", "ครุภัณฑ์", "ราคากลาง", "สัญญา"),
    "sso": ("ประกันสังคม", "ผู้ประกันตน", "กองทุนประกันสังคม"),
}


def number(value: object) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return 0.0


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("inventory", type=pathlib.Path)
    parser.add_argument("download_dir", type=pathlib.Path)
    parser.add_argument("output", type=pathlib.Path)
    args = parser.parse_args()

    inventory = json.loads(args.inventory.read_text(encoding="utf-8"))
    files = sorted(
        (item for item in inventory["files"] if item["category"] == "PBO" and item["title"].endswith(".xlsx")),
        key=lambda item: item["title"],
    )
    results = []
    for item in files:
        year = int(pathlib.Path(item["title"]).stem)
        path = args.download_dir / f'{item["id"]}.xlsx'
        with path.open("rb") as handle:
            workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
            data_sheet = max(workbook.worksheets, key=lambda sheet: sheet.max_row)
            totals = defaultdict(float)
            theme_totals = {key: defaultdict(float) for key in THEMES}
            theme_rows = defaultdict(int)
            row_count = 0
            rows = data_sheet.iter_rows(min_col=1, max_col=22, values_only=True)
            next(rows)
            first_data_row = next(rows)
            if first_data_row[0] == "Grand Total":
                totals["act"] = number(first_data_row[10])
                totals["adjusted"] = number(first_data_row[11])
                totals["paid"] = number(first_data_row[13])
                totals["committed"] = number(first_data_row[14])
                row_count = max(0, data_sheet.max_row - 2)
            else:
                for values in chain((first_data_row,), rows):
                    row_count += 1
                    act = number(values[10])
                    adjusted = number(values[11])
                    paid = number(values[13])
                    committed = number(values[14])
                    totals["act"] += act
                    totals["adjusted"] += adjusted
                    totals["paid"] += paid
                    totals["committed"] += committed
                    text = " ".join(str(value) for value in values[3:10] if value not in (None, "-", "")).lower()
                    for key, words in THEMES.items():
                        if any(word in text for word in words):
                            theme_rows[key] += 1
                            theme_totals[key]["act"] += act
                            theme_totals[key]["adjusted"] += adjusted
                            theme_totals[key]["paid"] += paid
            workbook.close()
        result = {
            "year": year,
            "rows": row_count,
            "act": round(totals["act"], 4),
            "adjusted": round(totals["adjusted"], 4),
            "paid": round(totals["paid"], 4),
            "committed": round(totals["committed"], 4),
            "paid_rate": round(totals["paid"] / totals["adjusted"] * 100, 1) if totals["adjusted"] else None,
            "themes": {
                key: {
                    "rows": theme_rows[key],
                    "act": round(theme_totals[key]["act"], 4),
                    "adjusted": round(theme_totals[key]["adjusted"], 4),
                    "paid": round(theme_totals[key]["paid"], 4),
                }
                for key in THEMES
            },
            "source_url": item["url"],
        }
        results.append(result)
        print(json.dumps({key: result[key] for key in ("year", "rows", "act", "adjusted", "paid_rate")}, ensure_ascii=False), flush=True)

    output = {"series": results, "years": len(results), "row_count": sum(item["rows"] for item in results)}
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"years": output["years"], "rows": output["row_count"]}), flush=True)


if __name__ == "__main__":
    main()
