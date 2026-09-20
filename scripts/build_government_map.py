"""Join Bureaucrazy Lab's public organisation index with PBO FY2568 data.

The output keeps source URLs and only uses organisation facts that are visible on
the public pages. Budget totals are calculated independently from the PBO workbook.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import time
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import requests
from bs4 import BeautifulSoup


BASE_URL = "https://www.bureaucrazylab.org"
KNOWN_TYPES = {
    "รัฐวิสาหกิจ": "stateEnterprise",
    "องค์การมหาชน": "publicOrganization",
    "หน่วยงานภายใต้การกำกับ": "other",
    "สถาบันอุดมศึกษา": "other",
    "อื่นๆ": "other",
}
PBO_MINISTRY_ALIASES = {
    "ส่วนราชการไม่สังกัดสำนักนายกรัฐมนตรีกระทรวงหรือทบวงและหน่วยงานภายใต้การควบคุมดูแลของนายกรัฐมนตรี":
        "ส่วนราชการไม่สังกัดสำนักนายกรัฐมนตรีกระทรวงหรือทบวงและหน่วยงานภายใต้การควบคุม",
}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def key(value: Any) -> str:
    text = unicodedata.normalize("NFC", clean(value)).replace("สํ", "สำ")
    text = text.replace("เเละ", "และ")
    return re.sub(r"[\s()\[\]{}.,/\\\-]+", "", text).lower()


def number(value: Any) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
        return result if math.isfinite(result) else 0.0
    return 0.0


def rounded(value: float) -> float:
    return round(value, 4)


def next_chunks(soup: BeautifulSoup) -> list[str]:
    chunks: list[str] = []
    prefix = "self.__next_f.push("
    for node in soup.find_all("script"):
        text = node.string or node.get_text()
        if not text.startswith(prefix):
            continue
        try:
            payload = json.loads(text[len(prefix) : -1])
        except json.JSONDecodeError:
            continue
        if len(payload) > 1 and isinstance(payload[1], str):
            chunks.append(payload[1])
    return chunks


def extract_largest_list(soup: BeautifulSoup, marker: str) -> list[dict[str, Any]]:
    candidates: list[list[dict[str, Any]]] = []
    decoder = json.JSONDecoder()
    for chunk in next_chunks(soup):
        start = chunk.find(marker)
        if start < 0:
            continue
        try:
            value = decoder.raw_decode(chunk[start + len(marker) :])[0]
        except (json.JSONDecodeError, ValueError):
            continue
        if isinstance(value, list) and all(isinstance(item, dict) for item in value):
            candidates.append(value)
    if not candidates:
        raise ValueError(f"Cannot find structured list for {marker}")
    return max(candidates, key=len)


def get_soup(session: requests.Session, path: str, attempts: int = 3) -> BeautifulSoup:
    url = f"{BASE_URL}{path}"
    last_error: Exception | None = None
    for attempt in range(attempts):
        try:
            response = session.get(url, timeout=45)
            response.raise_for_status()
            return BeautifulSoup(response.text, "html.parser")
        except requests.RequestException as error:
            last_error = error
            if attempt + 1 < attempts:
                time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"Unable to fetch {url}: {last_error}")


def parse_department(anchor: Any) -> dict[str, Any] | None:
    href = clean(anchor.get("href"))
    match = re.fullmatch(r"/department/(\d+)", href)
    heading = anchor.find("h3")
    if not match or not heading:
        return None
    description = clean(anchor.find("p").get_text(" ", strip=True)) if anchor.find("p") else ""
    badges = [clean(node.get_text(" ", strip=True)) for node in anchor.find_all("span")]
    department_type = "regular"
    for badge in badges:
        if badge in KNOWN_TYPES:
            department_type = KNOWN_TYPES[badge]
            break
    division_preview = [
        badge for badge in badges
        if badge not in KNOWN_TYPES and not re.fullmatch(r"\+\s*\d+\s*เพิ่มเติม", badge)
    ]
    extra = 0
    for badge in badges:
        extra_match = re.fullmatch(r"\+\s*(\d+)\s*เพิ่มเติม", badge)
        if extra_match:
            extra = int(extra_match.group(1))
    division_count = len(division_preview) + extra if department_type == "regular" else 0
    return {
        "id": int(match.group(1)),
        "name": clean(heading.get_text(" ", strip=True)),
        "description": description,
        "type": department_type,
        "divisionCount": division_count,
        "divisionPreview": division_preview,
        "sourceUrl": f"{BASE_URL}{href}",
    }


def pbo_rollups(workbook_path: Path) -> tuple[dict[str, dict[str, float]], dict[tuple[str, str], dict[str, float]]]:
    ministry = defaultdict(lambda: defaultdict(float))
    agency = defaultdict(lambda: defaultdict(float))
    with workbook_path.open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        sheet = max(workbook.worksheets, key=lambda item: item.max_row)
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        columns = {str(value): index for index, value in enumerate(headers)}
        required = {
            "ministry": columns["กระทรวง"],
            "agency": columns["หน่วยงาน"],
            "act": columns["พรบ. (ล้านบาท)"],
            "adjusted": columns["งบฯ หลังโอน/ปป. ทั้งสิ้น (ล้านบาท)"],
            "po": columns["PO ทั้งสิ้น (ล้านบาท)"],
            "paid": columns["เบิกจ่ายทั้งสิ้น (ล้านบาท)"],
            "paid_po": columns["เบิกจ่ายรวม PO (ล้านบาท)"],
        }
        for row in rows:
            ministry_name = clean(row[required["ministry"]])
            agency_name = clean(row[required["agency"]])
            act = number(row[required["act"]])
            adjusted = number(row[required["adjusted"]])
            committed = max(
                number(row[required["paid_po"]]),
                number(row[required["paid"]]) + number(row[required["po"]]),
            )
            for target in (ministry[key(ministry_name)], agency[(key(ministry_name), key(agency_name))]):
                target["rows"] += 1
                target["act"] += act
                target["adjusted"] += adjusted
                target["committed"] += committed
        workbook.close()
    return dict(ministry), dict(agency)


def candidates_by_scope(path: Path) -> tuple[dict[str, dict[str, float]], dict[tuple[str, str], dict[str, float]]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    ministry = defaultdict(lambda: defaultdict(float))
    agency = defaultdict(lambda: defaultdict(float))
    for item in data.get("items", []):
        amount = number(item.get("adjusted"))
        scopes = (
            ministry[key(item.get("ministry"))],
            agency[(key(item.get("ministry")), key(item.get("agency")))],
        )
        for target in scopes:
            target["candidateCount"] += 1
            target["candidateAmount"] += amount
    return dict(ministry), dict(agency)


def budget_payload(values: dict[str, float] | None, flags: dict[str, float] | None) -> dict[str, Any] | None:
    if not values:
        return None
    adjusted = values.get("adjusted", 0.0)
    committed = values.get("committed", 0.0)
    return {
        "rows": int(values.get("rows", 0)),
        "act": rounded(values.get("act", 0.0)),
        "adjusted": rounded(adjusted),
        "committed": rounded(committed),
        "rate": round(committed * 100 / adjusted, 1) if adjusted else None,
        "candidateCount": int((flags or {}).get("candidateCount", 0)),
        "candidateAmount": rounded((flags or {}).get("candidateAmount", 0.0)),
    }


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    parser.add_argument("--pbo", type=Path, required=True)
    parser.add_argument("--anomalies", type=Path, required=True)
    args = parser.parse_args()

    session = requests.Session()
    session.headers.update({"User-Agent": "Ngob-Gae public-interest research index/1.0"})
    home = get_soup(session, "/")
    ministries = extract_largest_list(home, '"ministries":')
    issues = extract_largest_list(get_soup(session, "/issues"), '"problems":')
    ministry_budget, agency_budget = pbo_rollups(args.pbo)
    ministry_flags, agency_flags = candidates_by_scope(args.anomalies)

    output_ministries = []
    linked_departments = 0
    for index, item in enumerate(ministries, 1):
        soup = get_soup(session, f"/ministry/{item['id']}")
        departments_by_id: dict[int, dict[str, Any]] = {}
        for anchor in soup.select('a[href^="/department/"]'):
            department = parse_department(anchor)
            if department:
                departments_by_id[department["id"]] = department
        departments = list(departments_by_id.values())
        ministry_name = clean(item.get("nameTh"))
        ministry_key = PBO_MINISTRY_ALIASES.get(key(ministry_name), key(ministry_name))
        for department in departments:
            lookup = (ministry_key, key(department["name"]))
            department["pbo2568"] = budget_payload(agency_budget.get(lookup), agency_flags.get(lookup))
            if department["pbo2568"]:
                linked_departments += 1
        departments.sort(key=lambda row: (-(row["pbo2568"] or {}).get("adjusted", -1), row["name"]))
        output_ministries.append({
            "id": item["id"],
            "name": ministry_name,
            "departmentCount": int(item.get("deptCount", len(departments))),
            "divisionCount": int(item.get("divisionCount", 0)),
            "totalSubunits": int(item.get("totalSubunits", 0)),
            "typeCounts": item.get("typeCounts", {}),
            "sourceUrl": f"{BASE_URL}/ministry/{item['id']}",
            "pbo2568": budget_payload(ministry_budget.get(ministry_key), ministry_flags.get(ministry_key)),
            "departments": departments,
        })
        print(f"[{index}/{len(ministries)}] {ministry_name}: {len(departments)} organisations")

    output_ministries.sort(key=lambda row: -((row["pbo2568"] or {}).get("adjusted", 0)))
    payload = {
        "meta": {
            "title": "แผนที่โครงสร้างรัฐเชื่อมงบประมาณ",
            "structureSource": "Bureaucrazy Lab",
            "structureSourceUrl": BASE_URL,
            "structureRetrieved": datetime.now().astimezone().isoformat(timespec="seconds"),
            "budgetSource": "PBO ผลการเบิกจ่ายงบประมาณปี 2568",
            "budgetSourceUrl": "https://drive.google.com/file/d/1f-lfPEsutobU6iB8W4LY1bThfhvYheHJ/view",
            "budgetYear": 2568,
            "unit": "ล้านบาท",
            "method": "เชื่อมชื่อกระทรวงและหน่วยงานหลังปรับ Unicode และช่องว่าง โดยไม่ใช้การเดาจากชื่อคล้าย",
            "ministries": len(output_ministries),
            "departments": sum(row["departmentCount"] for row in output_ministries),
            "divisions": sum(row["divisionCount"] for row in output_ministries),
            "linkedDepartments": linked_departments,
        },
        "issues": [
            {
                "slug": issue.get("slug"),
                "title": issue.get("titleTh"),
                "stats": issue.get("stats", {}),
                "sourceUrl": f"{BASE_URL}/issue/{issue.get('slug')}",
            }
            for issue in issues
        ],
        "ministries": output_ministries,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
