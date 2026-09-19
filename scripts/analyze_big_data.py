"""Build an explainable anomaly index from the PBO FY2568 line-item workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


GENERIC = re.compile(
    r"^(?:รายการ)?งบ(?:ประจำ|ดำเนินงาน|ลงทุน)(?:\s*\([^)]*\))?$|"
    r"^ค่าใช้จ่าย(?:อื่น|ในการดำเนินงาน)$|^ค่าใช้สอยอื่น$|"
    r"^รายการค่าใช้จ่าย$|^เงินอุดหนุนทั่วไป$",
    re.I,
)
CODE_ONLY = re.compile(r"^(?:[A-Za-zก-ฮ]{1,8}\.)?\s*\d+(?:[./-]\d+)*$")

THEMES = {
    "ict": re.compile(r"คอมพิวเตอร์|ซอฟต์แวร์|server|สารสนเทศ|ดิจิทัล|เครือข่าย|ฐานข้อมูล|cloud|cyber|ปัญญาประดิษฐ์|\bai\b", re.I),
    "training": re.compile(r"ฝึกอบรม|อบรม|สัมมนา|ศึกษาดูงาน|พัฒนาบุคลากร|ประชุมเชิงปฏิบัติการ", re.I),
    "construction": re.compile(r"ที่ดิน|สิ่งก่อสร้าง|ก่อสร้าง|ปรับปรุงอาคาร|อาคาร|ถนน|สะพาน|ระบายน้ำ", re.I),
    "consulting": re.compile(r"ที่ปรึกษา|จ้างศึกษา|ศึกษาวิจัย|วิจัยและพัฒนา", re.I),
    "procurement": re.compile(r"จัดซื้อ|จัดจ้าง|ประกวดราคา|พัสดุ|ราคากลาง|e-bidding", re.I),
    "ai": re.compile(r"ปัญญาประดิษฐ์|artificial intelligence|\bai\b|machine learning|deep learning|chatbot|แชตบอต|ระบบอัจฉริยะ", re.I),
    "sso": re.compile(r"สำนักงานประกันสังคม|กองทุนประกันสังคม|ผู้ประกันตน|เงินสมทบ", re.I),
}


def clean(value: Any) -> str:
    text = str(value or "").replace("\u2013", "-").replace("\u2014", "-")
    return re.sub(r"\s+", " ", text.strip())


def number(value: Any) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
        return result if math.isfinite(result) else 0.0
    return 0.0


def numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def percentile(values: list[float], position: float) -> float:
    if not values:
        return 0.0
    index = (len(values) - 1) * position
    lower = math.floor(index)
    upper = math.ceil(index)
    if lower == upper:
        return values[lower]
    return values[lower] + (values[upper] - values[lower]) * (index - lower)


def gini(values: list[float]) -> float:
    if not values or sum(values) == 0:
        return 0.0
    weighted = sum((index + 1) * value for index, value in enumerate(values))
    return (2 * weighted) / (len(values) * sum(values)) - (len(values) + 1) / len(values)


def rounded(value: float) -> float:
    return round(value, 4)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    with args.input.open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        sheet = max(workbook.worksheets, key=lambda item: item.max_row)
        rows = sheet.iter_rows(min_col=1, max_col=22, values_only=True)
        headers = next(rows)
        index = {str(header): position for position, header in enumerate(headers)}

        col = {
            "ministry": index["กระทรวง"],
            "agency": index["หน่วยงาน"],
            "plan": index["แผนงาน"],
            "project": index["งาน/โครงการ"],
            "item": index["ชื่อรหัสงบประมาณ"],
            "act": index["พรบ. (ล้านบาท)"],
            "adjusted": index["งบฯ หลังโอน/ปป. ทั้งสิ้น (ล้านบาท)"],
            "po": index["PO ทั้งสิ้น (ล้านบาท)"],
            "paid": index["เบิกจ่ายทั้งสิ้น (ล้านบาท)"],
            "paid_po": index["เบิกจ่ายรวม PO (ล้านบาท)"],
        }

        flag_stats = defaultdict(lambda: {"count": 0, "amount": 0.0})
        positive_amounts: list[float] = []
        candidates: list[dict[str, Any]] = []
        agencies = defaultdict(lambda: {"rows": 0, "adjusted": 0.0, "committed": 0.0})
        repeated = defaultdict(lambda: {"count": 0, "amount": 0.0, "agencies": set()})
        theme_stats = defaultdict(lambda: {"rows": 0, "adjusted": 0.0, "committed": 0.0})
        totals = defaultdict(float)
        threshold_counts = Counter()
        row_count = 0

        for values in rows:
            row_count += 1
            ministry = clean(values[col["ministry"]])
            agency = clean(values[col["agency"]])
            plan = clean(values[col["plan"]])
            project = clean(values[col["project"]])
            item = clean(values[col["item"]])
            act = number(values[col["act"]])
            adjusted = number(values[col["adjusted"]])
            paid = number(values[col["paid"]])
            po = number(values[col["po"]])
            paid_po = number(values[col["paid_po"]])
            execution_present = any(numeric(values[col[name]]) for name in ("po", "paid", "paid_po"))
            committed = max(paid_po, paid + po)
            rate = committed * 100 / adjusted if adjusted > 0 and execution_present else None
            delta = adjusted - act

            totals["act"] += act
            totals["adjusted"] += adjusted
            totals["paid"] += paid
            totals["committed"] += committed
            if adjusted > 0:
                positive_amounts.append(adjusted)
                if adjusted >= 20:
                    threshold_counts["20"] += 1
                if adjusted >= 100:
                    threshold_counts["100"] += 1
                if adjusted >= 500:
                    threshold_counts["500"] += 1

            agency_key = (ministry, agency)
            agencies[agency_key]["rows"] += 1
            agencies[agency_key]["adjusted"] += adjusted
            agencies[agency_key]["committed"] += committed

            normalized_item = re.sub(r"\d+", "#", item.lower())
            if adjusted > 0 and len(normalized_item) >= 8:
                repeated[normalized_item]["count"] += 1
                repeated[normalized_item]["amount"] += adjusted
                repeated[normalized_item]["agencies"].add(agency)

            haystack = " ".join((item, project, plan, agency))
            for theme, pattern in THEMES.items():
                if pattern.search(haystack):
                    theme_stats[theme]["rows"] += 1
                    theme_stats[theme]["adjusted"] += adjusted
                    theme_stats[theme]["committed"] += committed

            if adjusted < 20:
                continue

            signals: list[str] = []
            if act <= 0 and adjusted >= 20:
                signals.append("new_after_act")
            elif delta >= 20 and delta / max(act, 1) >= 0.25:
                signals.append("transfer_up")
            if act >= 20 and -delta >= 20 and -delta / act >= 0.25:
                signals.append("transfer_down")
            if adjusted >= 50 and not execution_present:
                signals.append("missing_execution")
            elif adjusted >= 50 and rate is not None and rate < 35:
                signals.append("low_execution")
            if adjusted >= 20 and rate is not None and rate > 105:
                signals.append("over_execution")
            vague = bool(GENERIC.search(item) or CODE_ONLY.fullmatch(item) or len(item) <= 7)
            if vague:
                signals.append("vague_title")
            if not signals:
                continue

            for signal in signals:
                flag_stats[signal]["count"] += 1
                flag_stats[signal]["amount"] += adjusted

            value_score = min(30.0, 10.0 + math.log10(max(adjusted / 20, 1)) * 10)
            movement_score = 0.0
            if "new_after_act" in signals:
                movement_score = 25.0
            elif "transfer_up" in signals or "transfer_down" in signals:
                movement_score = min(25.0, abs(delta) / max(act, adjusted, 1) * 50)
            execution_score = 20.0 if "missing_execution" in signals else 0.0
            if "low_execution" in signals and rate is not None:
                execution_score = min(25.0, (35 - rate) / 35 * 25 + 8)
            if "over_execution" in signals:
                execution_score = max(execution_score, 18.0)
            clarity_score = 20.0 if vague else 0.0
            score = min(100, round(value_score + movement_score + execution_score + clarity_score))
            identity = hashlib.sha1(f"{agency}|{item}|{adjusted}|{act}".encode("utf-8")).hexdigest()[:12]
            candidates.append({
                "id": identity,
                "score": score,
                "signals": signals,
                "ministry": ministry,
                "agency": agency,
                "plan": plan,
                "project": project,
                "item": item,
                "act": rounded(act),
                "adjusted": rounded(adjusted),
                "delta": rounded(delta),
                "committed": rounded(committed),
                "rate": round(rate, 1) if rate is not None else None,
                "execution_present": execution_present,
            })

        workbook.close()

    positive_amounts.sort()
    total_adjusted = sum(positive_amounts)
    top_one_count = max(1, math.ceil(len(positive_amounts) * 0.01))
    agency_rows = []
    for (ministry, agency), values in agencies.items():
        if values["adjusted"] <= 0:
            continue
        agency_rows.append({
            "ministry": ministry,
            "agency": agency,
            "rows": values["rows"],
            "adjusted": rounded(values["adjusted"]),
            "committed": rounded(values["committed"]),
            "rate": round(values["committed"] * 100 / values["adjusted"], 1),
            "share": round(values["adjusted"] * 100 / totals["adjusted"], 2),
        })

    repeated_rows = [
        {
            "pattern": pattern,
            "count": values["count"],
            "agencies": len(values["agencies"]),
            "adjusted": rounded(values["amount"]),
        }
        for pattern, values in repeated.items()
        if values["count"] >= 10 and len(values["agencies"]) >= 2
    ]

    labels = {
        "new_after_act": "วงเงินเกิดใหม่หลัง พ.ร.บ.",
        "transfer_up": "โอนเพิ่มอย่างน้อย 25%",
        "transfer_down": "ลดวงเงินอย่างน้อย 25%",
        "missing_execution": "ไม่มีตัวเลขใช้จ่ายระดับแถว",
        "low_execution": "ใช้จ่ายต่ำกว่า 35%",
        "over_execution": "ยอดรายงานเกินหลังโอน 5%",
        "vague_title": "ชื่อรายการกว้างหรือเป็นรหัส",
    }
    definitions = {
        "new_after_act": "วงเงินตาม พ.ร.บ. เป็นศูนย์หรือว่าง และวงเงินหลังโอนตั้งแต่ 20 ล้านบาท",
        "transfer_up": "วงเงินเพิ่มอย่างน้อย 20 ล้านบาทและไม่น้อยกว่า 25% ของวงเงินตั้งต้น",
        "transfer_down": "วงเงินลดอย่างน้อย 20 ล้านบาทและไม่น้อยกว่า 25% ของวงเงินตั้งต้น",
        "missing_execution": "วงเงินหลังโอนตั้งแต่ 50 ล้านบาท แต่ PO และยอดเบิกจ่ายระดับแถวไม่มีค่าตัวเลข",
        "low_execution": "วงเงินหลังโอนตั้งแต่ 50 ล้านบาท มีข้อมูลใช้จ่าย และอัตรารวม PO ต่ำกว่า 35%",
        "over_execution": "ยอดที่รายงานรวมเบิกจ่ายและ PO สูงกว่าวงเงินหลังโอนเกิน 5% ใช้เป็นจุดกระทบยอดนิยามและช่วงเวลา",
        "vague_title": "รายการตั้งแต่ 20 ล้านบาทที่ชื่อเป็นหมวดกว้าง รหัสสั้น หรือข้อความสั้นมาก",
    }
    flags = [
        {
            "id": key,
            "label": labels[key],
            "definition": definitions[key],
            "count": flag_stats[key]["count"],
            "amount": rounded(flag_stats[key]["amount"]),
        }
        for key in labels
    ]
    ranked_candidates = sorted(candidates, key=lambda item: (-item["score"], -item["adjusted"], item["item"]))
    output = {
        "meta": {
            "source": "PBO ผลการเบิกจ่ายงบประมาณปี 2568",
            "source_url": "https://drive.google.com/file/d/1f-lfPEsutobU6iB8W4LY1bThfhvYheHJ/view",
            "rows": row_count,
            "candidate_count": len(ranked_candidates),
            "unit": "ล้านบาท",
            "generated": "2569-09-19",
            "methodology_version": "1.1",
        },
        "overview": {
            "act": rounded(totals["act"]),
            "adjusted": rounded(totals["adjusted"]),
            "paid": rounded(totals["paid"]),
            "committed": rounded(totals["committed"]),
            "positive_rows": len(positive_amounts),
            "rows_over_20m": threshold_counts["20"],
            "rows_over_100m": threshold_counts["100"],
            "rows_over_500m": threshold_counts["500"],
            "median": rounded(statistics.median(positive_amounts)),
            "p95": rounded(percentile(positive_amounts, 0.95)),
            "p99": rounded(percentile(positive_amounts, 0.99)),
            "top_one_percent_share": round(sum(positive_amounts[-top_one_count:]) * 100 / total_adjusted, 1),
            "gini": round(gini(positive_amounts), 3),
        },
        "flags": flags,
        "items": ranked_candidates,
        "agencies": sorted(agency_rows, key=lambda item: -item["adjusted"])[:30],
        "themes": {
            key: {
                "rows": values["rows"],
                "adjusted": rounded(values["adjusted"]),
                "committed": rounded(values["committed"]),
                "rate": round(values["committed"] * 100 / values["adjusted"], 1) if values["adjusted"] else None,
            }
            for key, values in theme_stats.items()
        },
        "repeated_patterns": sorted(repeated_rows, key=lambda item: (-item["adjusted"], -item["count"]))[:30],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({
        "rows": row_count,
        "candidates": len(candidates),
        "published_items": len(output["items"]),
        "overview": output["overview"],
        "flags": flags,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
