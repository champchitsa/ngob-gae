"""Inspect every downloaded workbook and produce a compact corpus map.

This intentionally samples workbook structure before deeper, source-specific
analysis. It records every sheet, dimensions, likely headers, keywords and the
largest numeric values visible near the top of each sheet.
"""

from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

import openpyxl
import xlrd


KEYWORDS = {
    "ict": ("คอมพิวเตอร์", "ซอฟต์แวร์", "ดิจิทัล", "สารสนเทศ", "ระบบ"),
    "training": ("ฝึกอบรม", "อบรม", "สัมมนา", "ดูงาน"),
    "land": ("ก่อสร้าง", "ปรับปรุง", "ที่ดิน", "อาคาร", "ถนน"),
    "procurement": ("จัดซื้อ", "จัดจ้าง", "ครุภัณฑ์", "ราคากลาง", "สัญญา"),
    "subsidy": ("เงินอุดหนุน", "อุดหนุน"),
}


def safe_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def summarize_sample(rows: list[list[object]]) -> dict:
    strings = [safe_text(value) for row in rows for value in row if value is not None]
    combined = " ".join(strings)
    numbers = [
        float(value)
        for row in rows
        for value in row
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    ]
    likely_headers = []
    for row in rows[:15]:
        text_cells = [safe_text(value) for value in row if safe_text(value)]
        if len(text_cells) >= 2:
            likely_headers.append(text_cells[:12])
    return {
        "keyword_hits": {
            key: sum(combined.lower().count(word.lower()) for word in words)
            for key, words in KEYWORDS.items()
        },
        "largest_sample_numbers": sorted(numbers, reverse=True)[:8],
        "likely_headers": likely_headers[:5],
    }


def inspect_xlsx(path: pathlib.Path) -> list[dict]:
    file_handle = path.open("rb")
    workbook = openpyxl.load_workbook(file_handle, read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 35:
                    break
                rows.append(list(row[:60]))
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    **summarize_sample(rows),
                }
            )
    finally:
        workbook.close()
        file_handle.close()
    return sheets


def inspect_xls(path: pathlib.Path) -> list[dict]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    sheets = []
    try:
        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            rows = [
                [sheet.cell_value(row, col) for col in range(min(sheet.ncols, 60))]
                for row in range(min(sheet.nrows, 35))
            ]
            sheets.append(
                {
                    "name": sheet_name,
                    "rows": sheet.nrows,
                    "columns": sheet.ncols,
                    **summarize_sample(rows),
                }
            )
            workbook.unload_sheet(sheet_name)
    finally:
        workbook.release_resources()
    return sheets


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("inventory", type=pathlib.Path)
    parser.add_argument("download_dir", type=pathlib.Path)
    parser.add_argument("output", type=pathlib.Path)
    args = parser.parse_args()

    inventory = json.loads(args.inventory.read_text(encoding="utf-8"))
    workbooks = [
        item for item in inventory["files"] if pathlib.Path(item["title"]).suffix.lower() in {".xlsx", ".xls"}
    ]
    results = []
    for index, item in enumerate(workbooks, start=1):
        suffix = pathlib.Path(item["title"]).suffix.lower()
        path = args.download_dir / f'{item["id"]}{suffix}'
        result = {
            "id": item["id"],
            "title": item["title"],
            "category": item["category"],
            "path": item["path"],
            "source_url": item["url"],
            "size": item["size"],
        }
        if not path.exists():
            result["error"] = "not downloaded"
        else:
            try:
                with path.open("rb") as signature_file:
                    is_zip_workbook = signature_file.read(4) == b"PK\x03\x04"
                result["sheets"] = inspect_xlsx(path) if suffix == ".xlsx" or is_zip_workbook else inspect_xls(path)
            except Exception as error:
                result["error"] = f"{type(error).__name__}: {error}"
        results.append(result)
        print(f'[{index}/{len(workbooks)}] {item["title"]}', flush=True)

    summary = {
        "workbook_count": len(workbooks),
        "readable": sum("sheets" in item for item in results),
        "errors": sum("error" in item for item in results),
        "sheet_count": sum(len(item.get("sheets", [])) for item in results),
        "items": results,
    }
    args.output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: summary[key] for key in ("workbook_count", "readable", "errors", "sheet_count")}))


if __name__ == "__main__":
    main()
