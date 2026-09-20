"""Extract every file in the public Drive inventory into a resumable corpus.

The extractor keeps one gzip-compressed JSONL file per Drive file. Each output
starts with a file record, continues with one record per PDF page, spreadsheet
row, document paragraph/table row, slide, or image, and ends with a summary.
Large source files are downloaded to a temporary directory and removed after a
successful or failed extraction so the whole 6.67 GB archive never has to fit
on disk at once.
"""

from __future__ import annotations

import argparse
import concurrent.futures
import contextlib
import gzip
import hashlib
import heapq
import io
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from typing import Iterable

import openpyxl
import xlrd
from docx import Document
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader


THEMES = {
    "ict": ("คอมพิวเตอร์", "ซอฟต์แวร์", "ดิจิทัล", "สารสนเทศ", "เครือข่าย", "ฐานข้อมูล", "cloud", "license"),
    "ai": ("ปัญญาประดิษฐ์", " ai ", "machine learning", "ระบบอัจฉริยะ"),
    "procurement": ("จัดซื้อ", "จัดจ้าง", "ราคากลาง", "tor", "สัญญา", "ประกวดราคา", "ผู้รับจ้าง"),
    "construction": ("ก่อสร้าง", "ปรับปรุง", "ซ่อมแซม", "อาคาร", "ถนน", "สะพาน", "ขุดลอก"),
    "training": ("ฝึกอบรม", "อบรม", "สัมมนา", "ศึกษาดูงาน"),
    "land": ("ที่ดิน", "เวนคืน", "กรรมสิทธิ์", "ค่าเช่า"),
    "sso": ("ประกันสังคม", "สำนักงานประกันสังคม", "กองทุนประกันสังคม"),
    "health": ("สาธารณสุข", "โรงพยาบาล", "สุขภาพ", "ผู้ป่วย", "ยา"),
    "education": ("การศึกษา", "โรงเรียน", "นักเรียน", "มหาวิทยาลัย", "ครู"),
    "energy": ("พลังงาน", "ไฟฟ้า", "แสงอาทิตย์", "เชื้อเพลิง"),
    "defence": ("กองทัพ", "กลาโหม", "อาวุธ", "กระสุน", "ความมั่นคง"),
    "local": ("กรุงเทพมหานคร", "เทศบาล", "องค์การบริหารส่วน", "อบจ.", "อบต."),
}

NUMBER_RE = re.compile(r"(?<![\w.])[-+]?\d[\d,]*(?:\.\d+)?")
SPACE_RE = re.compile(r"[ \t\u00a0]+")
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_P = f"{{{WORD_NS}}}p"
W_T = f"{{{WORD_NS}}}t"
W_TAB = f"{{{WORD_NS}}}tab"
W_BR = f"{{{WORD_NS}}}br"
W_CR = f"{{{WORD_NS}}}cr"
W_TBL = f"{{{WORD_NS}}}tbl"
W_TR = f"{{{WORD_NS}}}tr"
W_TC = f"{{{WORD_NS}}}tc"
W_TXBX_CONTENT = f"{{{WORD_NS}}}txbxContent"
W_ID = f"{{{WORD_NS}}}id"
W_TYPE = f"{{{WORD_NS}}}type"


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(SPACE_RE.sub(" ", line).strip() for line in text.split("\n")).strip()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_text(value)


def text_is_sufficient(text: str) -> bool:
    useful = re.sub(r"[^\w\u0E00-\u0E7F]", "", text)
    return len(useful) >= 60


class FileStats:
    def __init__(self) -> None:
        self.units = 0
        self.lines = 0
        self.characters = 0
        self.nonempty_units = 0
        self.empty_units = 0
        self.ocr_units = 0
        self.embedded_units = 0
        self.cells = 0
        self.keyword_hits: Counter[str] = Counter()
        self._amounts: list[tuple[float, str]] = []
        self._sha = hashlib.sha256()

    def add(self, text: str, method: str, cells: int = 0) -> None:
        self.units += 1
        self.cells += cells
        self.characters += len(text)
        self.lines += len(text.splitlines()) if text else 0
        if text.strip():
            self.nonempty_units += 1
        else:
            self.empty_units += 1
        if method == "ocr":
            self.ocr_units += 1
        elif method == "embedded":
            self.embedded_units += 1
        lowered = f" {text.lower()} "
        for theme, words in THEMES.items():
            self.keyword_hits[theme] += sum(lowered.count(word) for word in words)
        self._sha.update(text.encode("utf-8", errors="replace"))
        for match in NUMBER_RE.finditer(text):
            raw = match.group(0).replace(",", "")
            with contextlib.suppress(ValueError):
                number = abs(float(raw))
                if number < 1000:
                    continue
                start = max(0, match.start() - 70)
                end = min(len(text), match.end() + 90)
                context = clean_text(text[start:end].replace("\n", " "))
                entry = (number, context[:220])
                if len(self._amounts) < 25:
                    heapq.heappush(self._amounts, entry)
                elif number > self._amounts[0][0]:
                    heapq.heapreplace(self._amounts, entry)

    def as_dict(self) -> dict:
        return {
            "units": self.units,
            "lines": self.lines,
            "characters": self.characters,
            "nonempty_units": self.nonempty_units,
            "empty_units": self.empty_units,
            "ocr_units": self.ocr_units,
            "embedded_units": self.embedded_units,
            "cells": self.cells,
            "keyword_hits": dict(self.keyword_hits),
            "largest_numbers": [
                {"value": value, "context": context}
                for value, context in sorted(self._amounts, reverse=True)
            ],
            "text_sha256": self._sha.hexdigest(),
        }


class JsonlWriter:
    def __init__(self, path: pathlib.Path) -> None:
        self.final_path = path
        self.temp_path = path.with_suffix(path.suffix + ".part")
        self.temp_path.parent.mkdir(parents=True, exist_ok=True)
        self.handle = gzip.open(self.temp_path, "wt", encoding="utf-8", compresslevel=6)

    def write(self, record: dict) -> None:
        self.handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")

    def commit(self) -> None:
        self.handle.close()
        self.temp_path.replace(self.final_path)

    def abort(self) -> None:
        with contextlib.suppress(Exception):
            self.handle.close()
        self.temp_path.unlink(missing_ok=True)


def download_file(item: dict, target: pathlib.Path, retries: int = 3) -> None:
    expected = int(item.get("size") or 0)
    if target.exists() and (not expected or target.stat().st_size == expected):
        return
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_suffix(target.suffix + ".part")
    url = "https://drive.usercontent.google.com/download?" + urllib.parse.urlencode(
        {"id": item["id"], "export": "download", "confirm": "t"}
    )
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        temp.unlink(missing_ok=True)
        try:
            request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(request, timeout=300) as response, temp.open("wb") as handle:
                while chunk := response.read(1024 * 1024):
                    handle.write(chunk)
            actual = temp.stat().st_size
            if expected and actual != expected:
                raise IOError(f"downloaded {actual} bytes, expected {expected}")
            temp.replace(target)
            return
        except Exception as error:
            last_error = error
            temp.unlink(missing_ok=True)
            if attempt < retries:
                time.sleep(attempt * 2)
    raise RuntimeError(f"download failed after {retries} attempts: {last_error}")


def find_binary(name: str, candidates: Iterable[pathlib.Path]) -> pathlib.Path:
    found = shutil.which(name)
    if found:
        return pathlib.Path(found)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"required binary not found: {name}")


def ocr_image(image_path: pathlib.Path, tesseract: pathlib.Path, tessdata: pathlib.Path) -> str:
    descriptor, prepared_name = tempfile.mkstemp(suffix=".png", dir=tessdata.parent)
    os.close(descriptor)
    prepared = pathlib.Path(prepared_name)
    try:
        with Image.open(image_path) as image:
            image = ImageOps.exif_transpose(image)
            if max(image.size) > 1400:
                scale = 1400 / max(image.size)
                image = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))), Image.Resampling.LANCZOS)
            image = ImageOps.autocontrast(image.convert("L"))
            image.save(prepared, format="PNG", optimize=True)
        command = [
            str(tesseract), str(prepared), "stdout", "--tessdata-dir", str(tessdata),
            "-l", "tha+eng", "--psm", "6", "-c", "preserve_interword_spaces=1",
        ]
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180)
        if result.returncode:
            raise RuntimeError(result.stderr.strip() or f"tesseract exited {result.returncode}")
        return clean_text(result.stdout)
    finally:
        prepared.unlink(missing_ok=True)


def extract_largest_page_image(page: object, directory: pathlib.Path) -> pathlib.Path | None:
    best: tuple[int, bytes, str] | None = None
    for image in page.images:
        try:
            payload = image.data
            with Image.open(io.BytesIO(payload)) as candidate:
                area = candidate.width * candidate.height
                image_format = (candidate.format or "png").lower()
            if best is None or area > best[0]:
                best = (area, payload, image_format)
        except Exception:
            continue
    if best is None:
        return None
    suffix = ".jpg" if best[2] in {"jpg", "jpeg"} else f".{best[2]}"
    target = directory / f"source{suffix}"
    target.write_bytes(best[1])
    return target


def render_pdf_pages(pdf_path: pathlib.Path, first_page: int, last_page: int, directory: pathlib.Path, pdftoppm: pathlib.Path) -> dict[int, pathlib.Path]:
    prefix = directory / "page"
    command = [
        str(pdftoppm), "-f", str(first_page), "-l", str(last_page), "-r", "145",
        "-png", str(pdf_path), str(prefix),
    ]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=max(240, (last_page - first_page + 1) * 45))
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or f"pdftoppm exited {result.returncode}")
    rendered = {}
    for image_path in directory.glob("page-*.png"):
        match = re.search(r"-(\d+)\.png$", image_path.name)
        if match:
            rendered[int(match.group(1))] = image_path
    if not rendered:
        raise RuntimeError("pdftoppm did not create page images")
    return rendered


def extract_pdf(path: pathlib.Path, writer: JsonlWriter, stats: FileStats, tools: dict) -> dict:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        with contextlib.suppress(Exception):
            reader.decrypt("")
    page_errors = []
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception as error:
            text = ""
            page_errors.append({"page": page_number, "stage": "text", "error": f"{type(error).__name__}: {error}"})
        pages.append({"page": page_number, "text": text, "method": "embedded"})

    missing = [entry["page"] for entry in pages if not text_is_sufficient(entry["text"])]
    render_fallback = []
    for page_number in missing:
        entry = pages[page_number - 1]
        with tempfile.TemporaryDirectory(dir=tools["temp_root"]) as directory:
            try:
                image_path = extract_largest_page_image(reader.pages[page_number - 1], pathlib.Path(directory))
                if image_path is None:
                    render_fallback.append(page_number)
                    continue
                ocr_text = ocr_image(image_path, tools["tesseract"], tools["tessdata"])
                if len(ocr_text) >= len(entry["text"]):
                    entry["text"] = ocr_text
                    entry["method"] = "ocr"
            except Exception as error:
                page_errors.append({"page": page_number, "stage": "direct_image_ocr", "error": f"{type(error).__name__}: {error}"})
                render_fallback.append(page_number)

    for batch_start in range(0, len(render_fallback), 20):
        batch_pages = render_fallback[batch_start:batch_start + 20]
        first_page, last_page = min(batch_pages), max(batch_pages)
        with tempfile.TemporaryDirectory(dir=tools["temp_root"]) as directory:
            try:
                rendered = render_pdf_pages(path, first_page, last_page, pathlib.Path(directory), tools["pdftoppm"])
            except Exception as error:
                for page_number in batch_pages:
                    page_errors.append({"page": page_number, "stage": "render", "error": f"{type(error).__name__}: {error}"})
                continue
            for page_number in batch_pages:
                entry = pages[page_number - 1]
                try:
                    image_path = rendered.get(page_number)
                    if not image_path:
                        raise RuntimeError("rendered page image was not found")
                    ocr_text = ocr_image(image_path, tools["tesseract"], tools["tessdata"])
                    if len(ocr_text) >= len(entry["text"]):
                        entry["text"] = ocr_text
                        entry["method"] = "ocr"
                except Exception as error:
                    page_errors.append({"page": page_number, "stage": "ocr", "error": f"{type(error).__name__}: {error}"})
                    entry["method"] = "embedded" if entry["text"] else "ocr_error"

    for entry in pages:
        page_number, text, method = entry["page"], entry["text"], entry["method"]
        stats.add(text, method)
        writer.write({"type": "page", "page": page_number, "method": method, "line_count": len(text.splitlines()), "text": text})
    return {"pages": len(reader.pages), "page_errors": page_errors}


def extract_xlsx(path: pathlib.Path, writer: JsonlWriter, stats: FileStats) -> dict:
    source = path.open("rb")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            nonempty_rows = 0
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if not any(values):
                    continue
                nonempty_rows += 1
                text = "\t".join(values)
                stats.add(text, "cell", sum(bool(value) for value in values))
                writer.write({"type": "row", "sheet": sheet.title, "row": row_number, "cells": values, "text": text})
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "nonempty_rows": nonempty_rows})
    finally:
        workbook.close()
        source.close()
    return {"sheets": sheets, "sheet_count": len(sheets)}


def extract_xls(path: pathlib.Path, writer: JsonlWriter, stats: FileStats) -> dict:
    workbook = xlrd.open_workbook(path, on_demand=True)
    sheets = []
    try:
        for name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(name)
            nonempty_rows = 0
            for row_number in range(sheet.nrows):
                values = [cell_text(sheet.cell_value(row_number, column)) for column in range(sheet.ncols)]
                while values and not values[-1]:
                    values.pop()
                if not any(values):
                    continue
                nonempty_rows += 1
                text = "\t".join(values)
                stats.add(text, "cell", sum(bool(value) for value in values))
                writer.write({"type": "row", "sheet": name, "row": row_number + 1, "cells": values, "text": text})
            sheets.append({"name": name, "rows": sheet.nrows, "columns": sheet.ncols, "nonempty_rows": nonempty_rows})
            workbook.unload_sheet(name)
    finally:
        workbook.release_resources()
    return {"sheets": sheets, "sheet_count": len(sheets)}


def ocr_embedded_blob(blob: bytes, suffix: str, tools: dict) -> str:
    descriptor, file_name = tempfile.mkstemp(suffix=suffix, dir=tools["temp_root"])
    os.close(descriptor)
    image_path = pathlib.Path(file_name)
    try:
        image_path.write_bytes(blob)
        return ocr_image(image_path, tools["tesseract"], tools["tessdata"])
    finally:
        image_path.unlink(missing_ok=True)


def word_paragraph_text(paragraph: object, *, exclude_nested_textboxes: bool = False) -> str:
    parts = []
    for node in paragraph.iter():
        if exclude_nested_textboxes and node.tag == W_T and has_word_ancestor(node, paragraph, {W_TXBX_CONTENT}):
            continue
        if node.tag == W_T:
            parts.append(node.text or "")
        elif node.tag == W_TAB:
            parts.append("\t")
        elif node.tag in {W_BR, W_CR}:
            parts.append("\n")
    return clean_text("".join(parts))


def word_block_text(element: object) -> str:
    return clean_text("\n".join(filter(None, (word_paragraph_text(paragraph) for paragraph in element.iter(W_P)))))


def has_word_ancestor(element: object, root: object, tags: set[str]) -> bool:
    parent = element.getparent() if hasattr(element, "getparent") else None
    while parent is not None and parent is not root:
        if parent.tag in tags:
            return True
        parent = parent.getparent() if hasattr(parent, "getparent") else None
    return False


def extract_word_story(
    element: object,
    story: str,
    writer: JsonlWriter,
    stats: FileStats,
    *,
    part: str | None = None,
    paragraphs: bool = True,
    tables: bool = True,
    textboxes: bool = True,
) -> dict:
    counts = {"paragraphs": 0, "table_rows": 0, "textboxes": 0}
    base = {"story": story}
    if part:
        base["part"] = part

    if paragraphs:
        for source_index, paragraph in enumerate(element.iter(W_P), start=1):
            if has_word_ancestor(paragraph, element, {W_TBL, W_TXBX_CONTENT}):
                continue
            text = word_paragraph_text(paragraph, exclude_nested_textboxes=True)
            if not text:
                continue
            counts["paragraphs"] += 1
            stats.add(text, "document")
            writer.write({"type": f"{story}_paragraph", **base, "paragraph": source_index, "text": text})

    if tables:
        table_index = 0
        for table in element.iter(W_TBL):
            if has_word_ancestor(table, element, {W_TBL, W_TXBX_CONTENT}):
                continue
            table_index += 1
            row_index = 0
            for row in table.iter(W_TR):
                if has_word_ancestor(row, table, {W_TBL}):
                    continue
                cells = [child for child in row if child.tag == W_TC]
                values = []
                for cell in cells:
                    cell_lines = [
                        word_paragraph_text(paragraph, exclude_nested_textboxes=True)
                        for paragraph in cell.iter(W_P)
                        if not has_word_ancestor(paragraph, cell, {W_TBL, W_TXBX_CONTENT})
                    ]
                    values.append(clean_text("\n".join(filter(None, cell_lines))))
                if not any(values):
                    continue
                row_index += 1
                counts["table_rows"] += 1
                text = "\t".join(values)
                stats.add(text, "document", sum(bool(value) for value in values))
                writer.write({"type": f"{story}_table_row", **base, "table": table_index, "row": row_index, "cells": values, "text": text})

    if textboxes:
        for textbox_index, textbox in enumerate(element.iter(W_TXBX_CONTENT), start=1):
            text = word_block_text(textbox)
            if not text:
                continue
            counts["textboxes"] += 1
            stats.add(text, "document")
            writer.write({"type": "textbox", **base, "textbox": textbox_index, "text": text})
    return counts


def extract_word_notes(path: pathlib.Path, writer: JsonlWriter, stats: FileStats) -> dict:
    counts = {"footnotes": 0, "endnotes": 0}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for note_type, member, element_name in (
            ("footnote", "word/footnotes.xml", "footnote"),
            ("endnote", "word/endnotes.xml", "endnote"),
        ):
            if member not in names:
                continue
            root = ET.fromstring(archive.read(member))
            for note in root.findall(f"{{{WORD_NS}}}{element_name}"):
                if note.attrib.get(W_TYPE) in {"separator", "continuationSeparator", "continuationNotice"}:
                    continue
                text = word_block_text(note)
                if not text:
                    continue
                counts[f"{note_type}s"] += 1
                stats.add(text, "document")
                writer.write({"type": note_type, "note_id": note.attrib.get(W_ID), "text": text})
    return counts


def extract_docx(path: pathlib.Path, writer: JsonlWriter, stats: FileStats, tools: dict) -> dict:
    document = Document(path)
    paragraphs = 0
    table_rows = 0
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = word_paragraph_text(paragraph._p, exclude_nested_textboxes=True)
        if not text:
            continue
        paragraphs += 1
        stats.add(text, "document")
        writer.write({"type": "paragraph", "paragraph": index, "style": paragraph.style.name if paragraph.style else "", "text": text})
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [clean_text(cell.text) for cell in row.cells]
            text = "\t".join(values)
            table_rows += 1
            stats.add(text, "document", sum(bool(value) for value in values))
            writer.write({"type": "table_row", "table": table_index, "row": row_index, "cells": values, "text": text})
    body_extra = extract_word_story(document.element.body, "body", writer, stats, paragraphs=False, tables=False)

    header_parts = 0
    footer_parts = 0
    header_paragraphs = 0
    footer_paragraphs = 0
    header_table_rows = 0
    footer_table_rows = 0
    header_textboxes = 0
    footer_textboxes = 0
    seen_story_parts = set()
    for relation in document.part.rels.values():
        relation_type = str(relation.reltype)
        if not (relation_type.endswith("/header") or relation_type.endswith("/footer")):
            continue
        part = relation.target_part
        part_name = str(part.partname)
        if part_name in seen_story_parts:
            continue
        seen_story_parts.add(part_name)
        story = "header" if relation_type.endswith("/header") else "footer"
        story_counts = extract_word_story(part.element, story, writer, stats, part=part_name)
        if story == "header":
            header_parts += 1
            header_paragraphs += story_counts["paragraphs"]
            header_table_rows += story_counts["table_rows"]
            header_textboxes += story_counts["textboxes"]
        else:
            footer_parts += 1
            footer_paragraphs += story_counts["paragraphs"]
            footer_table_rows += story_counts["table_rows"]
            footer_textboxes += story_counts["textboxes"]

    comments = 0
    for comment in document.comments:
        text = word_block_text(comment._element)
        if not text:
            continue
        comments += 1
        stats.add(text, "document")
        writer.write({
            "type": "comment",
            "comment_id": comment.comment_id,
            "author": comment.author,
            "initials": comment.initials,
            "timestamp": comment.timestamp.isoformat() if comment.timestamp else None,
            "text": text,
        })
    note_counts = extract_word_notes(path, writer, stats)

    images = 0
    image_errors = []
    seen_image_parts = set()
    for part in document.part.package.parts:
        if not str(getattr(part, "content_type", "")).startswith("image/"):
            continue
        part_name = str(part.partname)
        if part_name in seen_image_parts:
            continue
        seen_image_parts.add(part_name)
        images += 1
        try:
            extension = pathlib.Path(part_name).suffix or ".png"
            text = ocr_embedded_blob(part.blob, extension, tools)
            stats.add(text, "ocr")
            writer.write({"type": "embedded_image", "image": images, "part": part_name, "method": "ocr", "text": text})
        except Exception as error:
            image_errors.append({"image": images, "part": part_name, "error": f"{type(error).__name__}: {error}"})
            stats.add("", "ocr")
            writer.write({"type": "embedded_image", "image": images, "part": part_name, "method": "ocr_error", "text": ""})
    return {
        "paragraphs": paragraphs,
        "tables": len(document.tables),
        "table_rows": table_rows,
        "textboxes": body_extra["textboxes"] + header_textboxes + footer_textboxes,
        "header_parts": header_parts,
        "header_paragraphs": header_paragraphs,
        "header_table_rows": header_table_rows,
        "footer_parts": footer_parts,
        "footer_paragraphs": footer_paragraphs,
        "footer_table_rows": footer_table_rows,
        "comments": comments,
        **note_counts,
        "images": images,
        "image_errors": image_errors,
    }


def iter_pptx_shapes(shapes: object, prefix: tuple[int, ...] = ()):
    for index, shape in enumerate(shapes, start=1):
        path = (*prefix, index)
        yield shape, path
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pptx_shapes(shape.shapes, path)


def extract_pptx(path: pathlib.Path, writer: JsonlWriter, stats: FileStats, tools: dict) -> dict:
    presentation = Presentation(path)
    images = 0
    text_shapes = 0
    table_rows = 0
    notes = 0
    image_errors = []
    image_cache = {}

    def shape_path_label(shape_path: tuple[int, ...]) -> str:
        return ".".join(str(value) for value in shape_path)

    def write_image(shape: object, slide_number: int, shape_path: tuple[int, ...], record_type: str, locator: str | None = None) -> None:
        nonlocal images
        if not hasattr(shape, "image"):
            return
        images += 1
        part_label = locator or shape_path_label(shape_path)
        try:
            blob = shape.image.blob
            digest = hashlib.sha256(blob).hexdigest()
            if digest not in image_cache:
                extension = f".{shape.image.ext or 'png'}"
                try:
                    image_cache[digest] = (ocr_embedded_blob(blob, extension, tools), None)
                except Exception as error:
                    image_cache[digest] = ("", f"{type(error).__name__}: {error}")
            image_text, cached_error = image_cache[digest]
            if cached_error:
                raise RuntimeError(cached_error)
            stats.add(image_text, "ocr")
            writer.write({"type": record_type, "slide": slide_number, "shape_path": part_label, "image": images, "method": "ocr", "text": image_text})
        except Exception as error:
            image_errors.append({"slide": slide_number, "shape_path": part_label, "image": images, "error": f"{type(error).__name__}: {error}"})
            stats.add("", "ocr")
            writer.write({"type": record_type, "slide": slide_number, "shape_path": part_label, "image": images, "method": "ocr_error", "text": ""})

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_texts = set()
        for shape, shape_path in iter_pptx_shapes(slide.shapes):
            path_label = shape_path_label(shape_path)
            if getattr(shape, "has_text_frame", False):
                text = clean_text(shape.text)
                if text:
                    text_shapes += 1
                    slide_texts.add(text)
                    stats.add(text, "slide")
                    writer.write({"type": "slide_text", "slide": slide_number, "shape_path": path_label, "text": text})
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows, start=1):
                    values = [clean_text(cell.text) for cell in row.cells]
                    if not any(values):
                        continue
                    table_rows += 1
                    text = "\t".join(values)
                    stats.add(text, "slide", sum(bool(value) for value in values))
                    writer.write({"type": "slide_table_row", "slide": slide_number, "shape_path": path_label, "row": row_index, "cells": values, "text": text})
            write_image(shape, slide_number, shape_path, "slide_image")

        if slide.has_notes_slide:
            note_text_frame = slide.notes_slide.notes_text_frame
            primary_note = clean_text(note_text_frame.text) if note_text_frame else ""
            if primary_note:
                notes += 1
                stats.add(primary_note, "slide")
                writer.write({"type": "slide_note", "slide": slide_number, "note": notes, "shape_path": "notes", "text": primary_note})
            for shape, shape_path in iter_pptx_shapes(slide.notes_slide.shapes):
                path_label = f"notes.{shape_path_label(shape_path)}"
                if getattr(shape, "has_text_frame", False):
                    text = clean_text(shape.text)
                    if text and text != primary_note and text not in slide_texts:
                        notes += 1
                        stats.add(text, "slide")
                        writer.write({"type": "slide_note", "slide": slide_number, "note": notes, "shape_path": path_label, "text": text})
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows, start=1):
                        values = [clean_text(cell.text) for cell in row.cells]
                        if not any(values):
                            continue
                        table_rows += 1
                        text = "\t".join(values)
                        stats.add(text, "slide", sum(bool(value) for value in values))
                        writer.write({"type": "slide_note_table_row", "slide": slide_number, "shape_path": path_label, "row": row_index, "cells": values, "text": text})
                write_image(shape, slide_number, shape_path, "slide_note_image", path_label)
    return {"slides": len(presentation.slides), "text_shapes": text_shapes, "table_rows": table_rows, "notes": notes, "images": images, "image_errors": image_errors}


def extract_image(path: pathlib.Path, writer: JsonlWriter, stats: FileStats, tools: dict) -> dict:
    with Image.open(path) as image:
        width, height = image.size
    text = ocr_image(path, tools["tesseract"], tools["tessdata"])
    stats.add(text, "ocr")
    writer.write({"type": "image", "method": "ocr", "width": width, "height": height, "text": text})
    return {"images": 1, "width": width, "height": height}


def extract_unknown(path: pathlib.Path, writer: JsonlWriter, stats: FileStats, tools: dict) -> dict:
    signature = path.read_bytes()[:8]
    if signature.startswith(b"%PDF"):
        return {"detected_as": "pdf", **extract_pdf(path, writer, stats, tools)}
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
        text = "\n".join(names)
        stats.add(text, "archive")
        writer.write({"type": "archive_index", "entries": names, "text": text})
        return {"detected_as": "zip", "entries": len(names)}
    text = clean_text(path.read_text(encoding="utf-8", errors="replace"))
    stats.add(text, "text")
    writer.write({"type": "text", "text": text})
    return {"detected_as": "text"}


def process_item(item: dict, args: argparse.Namespace, tools: dict) -> dict:
    started = time.time()
    suffix = pathlib.Path(item["title"]).suffix.lower() or ".bin"
    output_path = args.output_dir / "records" / f'{item["id"]}.jsonl.gz'
    if output_path.exists() and not args.force:
        cached_summary = None
        with gzip.open(output_path, "rt", encoding="utf-8") as handle:
            for line in handle:
                record = json.loads(line)
                if record.get("type") == "summary":
                    cached_summary = record
        if cached_summary and cached_summary.get("id") == item["id"]:
            cached_summary.pop("type", None)
            return {**cached_summary, "status": "cached", "output": str(output_path), "output_bytes": output_path.stat().st_size}
        output_path.unlink(missing_ok=True)

    existing_machine = args.machine_dir / f'{item["id"]}{suffix}'
    owns_source = not existing_machine.exists()
    source_path = existing_machine if existing_machine.exists() else args.output_dir / "temp" / f'{item["id"]}{suffix}'
    writer = JsonlWriter(output_path)
    try:
        if owns_source:
            download_file(item, source_path)
        writer.write({
            "type": "file", "id": item["id"], "title": item["title"], "mime_type": item.get("type"),
            "category": item.get("category"), "path": item.get("path"), "source_url": item.get("url"),
            "source_bytes": item.get("size"),
        })
        stats = FileStats()
        signature = source_path.read_bytes()[:8]
        if suffix == ".pdf" or signature.startswith(b"%PDF"):
            structure = extract_pdf(source_path, writer, stats, tools)
        elif suffix == ".xlsx" or (signature.startswith(b"PK") and (suffix == ".xls" or item.get("type", "").endswith("spreadsheetml.sheet"))):
            structure = extract_xlsx(source_path, writer, stats)
        elif suffix == ".xls":
            structure = extract_xls(source_path, writer, stats)
        elif suffix == ".docx":
            structure = extract_docx(source_path, writer, stats, tools)
        elif suffix == ".pptx":
            structure = extract_pptx(source_path, writer, stats, tools)
        elif suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"}:
            structure = extract_image(source_path, writer, stats, tools)
        else:
            structure = extract_unknown(source_path, writer, stats, tools)
        summary = {
            "id": item["id"], "title": item["title"], "status": "complete", "output": str(output_path),
            "source_bytes": item.get("size"), "duration_seconds": round(time.time() - started, 2),
            "structure": structure, **stats.as_dict(),
        }
        writer.write({"type": "summary", **summary})
        writer.commit()
        summary["output_bytes"] = output_path.stat().st_size
        return summary
    except Exception as error:
        writer.abort()
        return {
            "id": item["id"], "title": item["title"], "status": "error",
            "source_bytes": item.get("size"), "duration_seconds": round(time.time() - started, 2),
            "error": f"{type(error).__name__}: {error}",
        }
    finally:
        if owns_source:
            source_path.unlink(missing_ok=True)
            source_path.with_suffix(source_path.suffix + ".part").unlink(missing_ok=True)


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("inventory", type=pathlib.Path)
    parser.add_argument("output_dir", type=pathlib.Path)
    parser.add_argument("--machine-dir", type=pathlib.Path, default=pathlib.Path(".workdata/drive-machine"))
    parser.add_argument("--workers", type=int, default=3)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--only", choices=["pdf", "sheet", "image", "document", "other"])
    parser.add_argument("--id", action="append", dest="ids", help="process only the selected Drive file id; may be repeated")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    inventory = json.loads(args.inventory.read_text(encoding="utf-8"))
    items = inventory["files"]
    if args.ids:
        selected_ids = set(args.ids)
        items = [item for item in items if item["id"] in selected_ids]
    if args.only:
        def item_group(item: dict) -> str:
            suffix = pathlib.Path(item["title"]).suffix.lower()
            mime_type = str(item.get("type") or "").lower()
            if suffix == ".pdf" or mime_type == "application/pdf":
                return "pdf"
            if suffix in {".xlsx", ".xls"} or "spreadsheet" in mime_type or mime_type == "application/vnd.ms-excel":
                return "sheet"
            if suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"} or mime_type.startswith("image/"):
                return "image"
            if suffix in {".docx", ".pptx"} or "wordprocessingml" in mime_type or "presentationml" in mime_type:
                return "document"
            return "other"

        items = [item for item in items if item_group(item) == args.only]
    if args.limit:
        items = items[: args.limit]

    runtime_root = pathlib.Path(os.environ.get("CODEX_RUNTIME_ROOT", pathlib.Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"))
    pdftoppm = find_binary("pdftoppm", [runtime_root / "native/poppler/Library/bin/pdftoppm.exe"])
    tesseract = find_binary("tesseract", [pathlib.Path("C:/Program Files/Tesseract-OCR/tesseract.exe")])
    tessdata = pathlib.Path(os.environ.get("NGOB_GAE_TESSDATA", pathlib.Path.home() / ".cache/ngob-gae-tessdata")).resolve()
    if not (tessdata / "tha.traineddata").exists() or not (tessdata / "eng.traineddata").exists():
        raise FileNotFoundError("Thai and English Tesseract language data are required")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "temp").mkdir(parents=True, exist_ok=True)
    ocr_temp_root = pathlib.Path(os.environ.get("NGOB_GAE_OCR_TEMP", pathlib.Path.home() / ".cache/ngob-gae-ocr-temp"))
    ocr_temp_root.mkdir(parents=True, exist_ok=True)
    tools = {"pdftoppm": pdftoppm, "tesseract": tesseract, "tessdata": tessdata, "temp_root": ocr_temp_root}
    manifest_path = args.output_dir / "manifest.json"
    manifest = {"started_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "inventory_count": len(inventory["files"]), "items": {}}
    if manifest_path.exists():
        with contextlib.suppress(Exception):
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest["inventory_count"] = len(inventory["files"])

    lock = threading.Lock()
    completed = 0
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        futures = {pool.submit(process_item, item, args, tools): item for item in items}
        for future in concurrent.futures.as_completed(futures):
            result = future.result()
            with lock:
                completed += 1
                manifest["items"][result["id"]] = result
                statuses = Counter(value.get("status") for value in manifest["items"].values())
                manifest["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
                manifest["status_counts"] = dict(statuses)
                temp_manifest = manifest_path.with_suffix(".json.part")
                temp_manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
                temp_manifest.replace(manifest_path)
            print(f'[{completed}/{len(items)}] {result["status"]}: {result["title"]}', flush=True)

    complete = [value for value in manifest["items"].values() if value.get("status") in {"complete", "cached"}]
    errors = [value for value in manifest["items"].values() if value.get("status") == "error"]
    print(json.dumps({"selected": len(items), "complete": len(complete), "errors": len(errors), "manifest": str(manifest_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
