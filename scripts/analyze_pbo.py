"""Build a compact, reproducible evidence pack from the PBO FY2568 workbook.

The script never labels a row as corrupt or unlawful. It prioritises records for
human review using transparent rules that can be reproduced from the source.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl


HEADERS = {
    "year": "ปีงบประมาณ",
    "ministry": "กระทรวง",
    "agency": "หน่วยงาน",
    "strategy": "ยุทธศาสตร์การจัดสรร",
    "plan": "แผนงาน",
    "output": "ผลผลิต/โครงการ",
    "project": "งาน/โครงการ",
    "expense": "งบรายจ่าย",
    "capital": "รายจ่ายประจำ/ลงทุน",
    "item": "ชื่อรหัสงบประมาณ",
    "act": "พรบ. (ล้านบาท)",
    "adjusted": "งบฯ หลังโอน/ปป. ทั้งสิ้น (ล้านบาท)",
    "po": "PO ทั้งสิ้น (ล้านบาท)",
    "paid": "เบิกจ่ายทั้งสิ้น (ล้านบาท)",
    "paid_po": "เบิกจ่ายรวม PO (ล้านบาท)",
    "reserved": "เงินกันฯ สุทธิ (ล้านบาท)",
    "carry_paid": "เบิกจ่ายเหลื่อมปี (ล้านบาท)",
    "remain_debt": "คงเหลือกรณีมีหนี้ผูกพัน (ล้านบาท)",
    "remain_free": "คงเหลือกรณีไม่มีหนี้ผูกพัน (ล้านบาท)",
    "remain_work": "คงเหลือ สรก.อยู่ระหว่างดำเนินการ (ล้านบาท)",
    "remain_extend": "คงเหลืออยู่ระหว่างกันและขยายรวม (ล้านบาท)",
    "remain_total": "คงเหลือรวม (ล้านบาท)",
}


THEMES = {
    "ict": re.compile(
        r"คอมพิวเตอร์|ซอฟต์แวร์|software|server|ระบบสารสนเทศ|ดิจิทัล|เครือข่าย|"
        r"อินเทอร์เน็ต|ฐานข้อมูล|cloud|cyber|ปัญญาประดิษฐ์|artificial intelligence|ai\b",
        re.I,
    ),
    "training": re.compile(r"ฝึกอบรม|อบรม|สัมมนา|ศึกษาดูงาน|พัฒนาบุคลากร|ประชุมเชิงปฏิบัติการ", re.I),
    "land": re.compile(r"ที่ดิน|สิ่งก่อสร้าง|ก่อสร้าง|ปรับปรุงอาคาร|อาคาร|ถนน|สะพาน|ระบบระบายน้ำ", re.I),
    "consulting": re.compile(r"ที่ปรึกษา|จ้างศึกษา|ศึกษาวิจัย|วิจัยและพัฒนา", re.I),
    "sso": re.compile(r"ประกันสังคม|ผู้ประกันตน|กองทุนเงินทดแทน", re.I),
}

GENERIC = re.compile(
    r"^(?:รายการ)?งบ(?:ประจำ|ดำเนินงาน|ลงทุน)|ค่าใช้จ่าย(?:อื่น|ในการดำเนินงาน)|"
    r"ค่าใช้สอยอื่น|รายการค่าใช้จ่าย|เงินอุดหนุนทั่วไป$",
    re.I,
)


def number(value: Any) -> float:
    if value is None or value == "-":
        return 0.0
    if isinstance(value, (int, float)):
        value = float(value)
        return value if math.isfinite(value) else 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0.0


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def round_money(value: float) -> float:
    return round(value, 4)


def pct(numerator: float, denominator: float) -> float | None:
    if denominator <= 0:
        return None
    return round(numerator * 100 / denominator, 1)


def money_row(row: tuple[Any, ...], index: dict[str, int]) -> dict[str, Any]:
    record = {key: row[index[label]] for key, label in HEADERS.items()}
    for key in (
        "act",
        "adjusted",
        "po",
        "paid",
        "paid_po",
        "reserved",
        "carry_paid",
        "remain_debt",
        "remain_free",
        "remain_work",
        "remain_extend",
        "remain_total",
    ):
        record[key] = number(record[key])
    for key in (
        "ministry",
        "agency",
        "strategy",
        "plan",
        "output",
        "project",
        "expense",
        "capital",
        "item",
    ):
        record[key] = clean(record[key])
    return record


def evidence_row(record: dict[str, Any]) -> dict[str, Any]:
    committed = max(record["paid_po"], record["paid"] + record["po"])
    return {
        "ministry": record["ministry"],
        "agency": record["agency"],
        "plan": record["plan"],
        "project": record["project"],
        "item": record["item"],
        "capital": record["capital"],
        "act": round_money(record["act"]),
        "adjusted": round_money(record["adjusted"]),
        "po": round_money(record["po"]),
        "paid": round_money(record["paid"]),
        "committed": round_money(committed),
        "remain_free": round_money(record["remain_free"]),
        "remain_total": round_money(record["remain_total"]),
        "rate": pct(committed, record["adjusted"]),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    sheet = max(workbook.worksheets, key=lambda s: s.max_row)
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    index = {label: i for i, label in enumerate(header)}
    missing = [label for label in HEADERS.values() if label not in index]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    totals = defaultdict(float)
    agency = defaultdict(lambda: defaultdict(float))
    theme = defaultdict(lambda: {"rows": 0, "adjusted": 0.0, "paid": 0.0, "committed": 0.0})
    low_absorption: list[dict[str, Any]] = []
    transfers_in: list[dict[str, Any]] = []
    opaque: list[dict[str, Any]] = []
    free_balance: list[dict[str, Any]] = []
    sso_rows: list[dict[str, Any]] = []
    row_count = 0

    for row in rows:
        record = money_row(row, index)
        row_count += 1
        committed = max(record["paid_po"], record["paid"] + record["po"])
        totals["act"] += record["act"]
        totals["adjusted"] += record["adjusted"]
        totals["paid"] += record["paid"]
        totals["po"] += record["po"]
        totals["committed"] += committed
        totals["remain_free"] += record["remain_free"]
        totals["remain_total"] += record["remain_total"]

        key = (record["ministry"], record["agency"])
        for name in ("act", "adjusted", "paid", "po", "remain_free", "remain_total"):
            agency[key][name] += record[name]
        agency[key]["committed"] += committed
        agency[key]["rows"] += 1

        haystack = " ".join(
            [record["item"], record["project"], record["output"], record["plan"], record["agency"]]
        )
        for name, pattern in THEMES.items():
            if pattern.search(haystack):
                theme[name]["rows"] += 1
                theme[name]["adjusted"] += record["adjusted"]
                theme[name]["paid"] += record["paid"]
                theme[name]["committed"] += committed

        adjusted = record["adjusted"]
        rate = committed / adjusted if adjusted else 1.0
        if adjusted >= 50 and rate < 0.35:
            low_absorption.append(evidence_row(record))

        transfer = adjusted - record["act"]
        if transfer >= 20 and (record["act"] == 0 or transfer / max(record["act"], 1) >= 0.25):
            item = evidence_row(record)
            item["transfer_in"] = round_money(transfer)
            transfers_in.append(item)

        if adjusted >= 20 and GENERIC.search(record["item"]):
            opaque.append(evidence_row(record))

        if record["remain_free"] >= 20:
            free_balance.append(evidence_row(record))

        if THEMES["sso"].search(haystack):
            sso_rows.append(evidence_row(record))

    agency_rollup = []
    for (ministry, name), values in agency.items():
        item = {"ministry": ministry, "agency": name, **{k: round_money(v) for k, v in values.items()}}
        item["rate"] = pct(values["committed"], values["adjusted"])
        agency_rollup.append(item)

    low_agencies = sorted(
        [x for x in agency_rollup if x["adjusted"] >= 500 and x["rate"] is not None],
        key=lambda x: (x["rate"], -x["adjusted"]),
    )[:20]

    for values in theme.values():
        values["rate"] = pct(values["committed"], values["adjusted"])
        for key in ("adjusted", "paid", "committed"):
            values[key] = round_money(values[key])

    report = {
        "meta": {
            "source": "PBO ผลการเบิกจ่ายงบประมาณปี 2568",
            "sheet": sheet.title,
            "rows": row_count,
            "unit": "ล้านบาท",
            "method_note": "committed = max(เบิกจ่ายรวม PO, เบิกจ่าย + PO); flags are review priorities, not findings of wrongdoing",
        },
        "totals": {k: round_money(v) for k, v in totals.items()},
        "themes": theme,
        "low_absorption_rows": sorted(low_absorption, key=lambda x: (x["rate"], -x["adjusted"]))[:40],
        "transfer_in_rows": sorted(transfers_in, key=lambda x: -x["transfer_in"])[:30],
        "opaque_rows": sorted(opaque, key=lambda x: -x["adjusted"])[:30],
        "free_balance_rows": sorted(free_balance, key=lambda x: -x["remain_free"])[:30],
        "low_absorption_agencies": low_agencies,
        "sso_rows": sorted(sso_rows, key=lambda x: -x["adjusted"])[:50],
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "rows": row_count, "totals": report["totals"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
